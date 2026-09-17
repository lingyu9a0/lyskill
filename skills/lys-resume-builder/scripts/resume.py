"""Case-based, evidence-linked resume assets and immutable JD branches."""
import argparse
from datetime import date, datetime, timezone
import difflib
import html
import json
from pathlib import Path
import re
import uuid


def require(ok, message):
    if not ok:
        raise ValueError(message)


def nonempty(value):
    return isinstance(value, str) and bool(value.strip())


def safe_name(value):
    require(nonempty(value), "Name required")
    text = re.sub(r"[\\/:*?\"<>|]+", "-", value.strip())
    text = re.sub(r"\s+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-.")
    require(bool(text), "Name contains no usable characters")
    return text[:80]


def timestamp_id():
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    return stamp + "-" + uuid.uuid4().hex[:12]


def atomic_write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + "." + uuid.uuid4().hex + ".tmp")
    temp.write_text(text)
    temp.replace(path)


def init_case(base, alias, case_date=None):
    alias_slug = safe_name(alias)
    day = case_date or date.today().isoformat()
    require(bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", day)), "Case date must be YYYY-MM-DD")
    case = base.expanduser().resolve() / (day + "-" + alias_slug)
    require(not case.exists(), "Case directory already exists; continue the existing case instead")
    (case / "jds").mkdir(parents=True)
    (case / "resumes").mkdir()
    (case / ".records").mkdir()
    (case / ".labels").mkdir()
    metadata = {
        "alias": alias,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "case_dir": str(case),
    }
    (case / ".case.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2))
    (case / "profile.md").write_text("# 个人资料\n\n尚待根据首次材料整理。\n")
    (case / "experience-master.md").write_text("# 真实经历母版\n\n尚待根据真实材料整理。\n")
    (case / "evidence.md").write_text("# 证据链\n\n尚待根据真实材料整理。\n")
    (case / "job-signals.md").write_text("# 岗位信号\n\n尚待从证据中提取。\n")
    (case / "changelog.md").write_text("# 变更记录\n")
    return {"case_dir": str(case), "alias": alias}


def require_case(case):
    case = case.expanduser().resolve()
    require((case / ".case.json").is_file(), "Not an initialized resume case")
    return case


def validate(data):
    require(data.get("kind") in ("master", "match", "tailored"), "Invalid kind")
    require(data.get("stage") in ("draft", "final"), "Invalid stage")
    require(data.get("kind") != "master" or data.get("stage") == "draft", "Master is an asset record, not a final resume")
    require(data.get("kind") != "match" or data.get("stage") == "draft", "Match record must stay draft")
    require(nonempty(data.get("change_reason")), "Change reason required")
    require(
        isinstance(data.get("changes"), list)
        and data["changes"]
        and all(nonempty(x) for x in data["changes"]),
        "Concrete change list required",
    )
    require(isinstance(data.get("profile"), dict), "Profile object required")
    require(isinstance(data.get("evidence"), list) and data["evidence"], "Real source material required")

    evidence = {}
    for item in data["evidence"]:
        require(all(nonempty(item.get(key)) for key in ("id", "fact", "source", "ownership")), "Evidence fields required")
        require(item["id"] not in evidence, "Duplicate evidence id")
        require(
            item.get("status") in ("材料已支持", "用户已确认", "材料冲突", "归属不清", "缺失"),
            "Invalid evidence status",
        )
        evidence[item["id"]] = item

    experiences = {}
    source_types = ("课程项目", "社团项目", "个人项目", "实习", "工作", "比赛", "兼职", "志愿活动", "其他")
    execution_statuses = ("已落地", "部分落地", "仅方案", "模拟")
    require(isinstance(data.get("experiences"), list) and data["experiences"], "Structured experiences required")
    for item in data["experiences"]:
        require(
            all(nonempty(item.get(key)) for key in ("experience_id", "name", "source_type", "execution_status", "start_date", "end_date", "duration")),
            "Experience source, status, and date fields required",
        )
        require(item["experience_id"] not in experiences, "Duplicate experience id")
        require(item["source_type"] in source_types, "Invalid experience source type")
        require(item["execution_status"] in execution_statuses, "Invalid experience execution status")
        refs = item.get("evidence_ids")
        require(isinstance(refs, list) and refs and all(ref in evidence for ref in refs), "Experience needs evidence references")
        experiences[item["experience_id"]] = item

    signals = data.get("job_signals")
    require(isinstance(signals, list), "Job signals must be a list")
    for signal in signals:
        require(nonempty(signal.get("name")) and nonempty(signal.get("scope")), "Invalid job signal")
        refs = signal.get("evidence_ids")
        require(
            isinstance(refs, list) and refs and all(ref in evidence for ref in refs),
            "Every job signal needs evidence references",
        )

    require(data.get("kind") == "master" or nonempty(data.get("target")), "Target required")
    if data.get("kind") == "master":
        require(data.get("workflow_state") in ("BUILDING_ASSETS", "WAITING_FOR_TARGET_JD"), "Master workflow state required")
    if data.get("kind") in ("match", "tailored"):
        require(nonempty(data.get("company")), "Company required for JD work")
        require(nonempty(data.get("jd_id")), "JD id required for JD work")
        require(nonempty(data.get("jd_text")), "JD required for JD work")
        clusters = data.get("task_clusters")
        require(isinstance(clusters, list) and 3 <= len(clusters) <= 6, "JD needs 3 to 6 core task clusters")
        task_ids = set()
        for cluster in clusters:
            require(
                all(nonempty(cluster.get(key)) for key in ("task_id", "name", "object", "action", "deliverable", "evidence_expected")),
                "Invalid task cluster",
            )
            require(cluster["task_id"] not in task_ids, "Duplicate task id")
            task_ids.add(cluster["task_id"])

        matrix = data.get("match_matrix")
        require(isinstance(matrix, list) and matrix, "JD work requires JD-signal-evidence matching")
        requirement_ids = set()
        candidates_by_task = {task_id: [] for task_id in task_ids}
        for item in matrix:
            require(
                nonempty(item.get("requirement_id"))
                and nonempty(item.get("requirement"))
                and item.get("task_id") in task_ids
                and isinstance(item.get("core"), bool)
                and (not item.get("experience_id") or item.get("experience_id") in experiences)
                and nonempty(item.get("action"))
                and nonempty(item.get("gap_resolution")),
                "Invalid match item",
            )
            require(item["requirement_id"] not in requirement_ids, "Duplicate requirement id")
            requirement_ids.add(item["requirement_id"])
            require(
                item.get("status") in ("已有直接证据", "有相邻／代理证据", "有经历但证据不足", "完全缺失"),
                "Invalid match status",
            )
            require(item.get("evidence_grade") in ("A", "B", "C", "D", "E"), "Invalid evidence grade")
            require(item.get("relevance") in ("核心相关", "高度相关", "部分相关", "弱相关"), "Invalid relevance level")
            require(item.get("contribution_clarity") in ("个人清晰", "团队边界清晰", "归属不清"), "Invalid contribution clarity")
            require(item.get("result_strength") in ("可观察结果", "明确交付物", "过程动作", "无结果"), "Invalid result strength")
            require(isinstance(item.get("concrete_task_action"), bool), "Concrete task action decision required")
            require(isinstance(item.get("selected_for_resume"), bool), "Evidence selection decision required")
            refs = item.get("evidence_ids")
            require(
                isinstance(refs, list) and all(ref in evidence for ref in refs),
                "Invalid match evidence references",
            )
            require(item["status"] not in ("已有直接证据", "有相邻／代理证据", "有经历但证据不足") or bool(refs), "Supported match needs evidence")
            if item["core"] and item["status"] != "已有直接证据":
                require(
                    item["gap_resolution"] in ("待回补", "用户确认无更多真实经历"),
                    "Core evidence gap must trigger backfill or record exhaustion",
                )
            if item["status"] == "已有直接证据":
                require(item["gap_resolution"] in ("不需要回补", "已回补"), "Direct evidence has invalid backfill state")
            if item["selected_for_resume"]:
                require(item["status"] in ("已有直接证据", "有相邻／代理证据"), "Unresolved evidence cannot be selected")
                require(item["evidence_grade"] != "E", "E-grade evidence stays in the asset library")
                require(item.get("experience_id") in experiences, "Selected evidence needs an experience id")
                require(item["contribution_clarity"] != "归属不清", "Unclear ownership cannot enter a resume")
                experience = experiences[item["experience_id"]]
                if experience["source_type"] == "课程项目" and item["evidence_grade"] in ("A", "B"):
                    require(item["concrete_task_action"] is True, "Course research needs concrete task action for A or B grade")
            candidates_by_task[item["task_id"]].append(item)

        grade_rank = {"A": 1, "B": 2, "C": 3, "D": 4, "E": 5}
        relevance_rank = {"核心相关": 1, "高度相关": 2, "部分相关": 3, "弱相关": 4}
        contribution_rank = {"个人清晰": 1, "团队边界清晰": 2, "归属不清": 3}
        result_rank = {"可观察结果": 1, "明确交付物": 2, "过程动作": 3, "无结果": 4}
        def selection_rank(item):
            return (
                relevance_rank[item["relevance"]],
                grade_rank[item["evidence_grade"]],
                contribution_rank[item["contribution_clarity"]],
                result_rank[item["result_strength"]],
            )
        for task_id, candidates in candidates_by_task.items():
            require(candidates, "Every task cluster needs an evidence mapping row")
            selected = [item for item in candidates if item["selected_for_resume"]]
            require(len(selected) <= 2, "Each task cluster may select at most two evidence groups")
            eligible = [item for item in candidates if item["status"] in ("已有直接证据", "有相邻／代理证据") and item["evidence_grade"] != "E"]
            if selected and eligible:
                require(
                    min(selection_rank(item) for item in selected) == min(selection_rank(item) for item in eligible),
                    "Selected evidence must follow relevance, directness, contribution, and result priority",
                )

        selected_rows = [item for item in matrix if item["selected_for_resume"]]
        selected_evidence = {ref for item in selected_rows for ref in item["evidence_ids"]}
        allowed = data.get("allowed_evidence_ids")
        require(isinstance(allowed, list) and len(allowed) == len(set(allowed)), "Allowed evidence set must be unique")
        require(set(allowed) == selected_evidence, "Allowed evidence set must exactly equal selected evidence")
        selected_experiences = {item["experience_id"] for item in selected_rows}
        order = data.get("experience_order")
        require(isinstance(order, list) and len(order) == len(set(order)), "Experience order must be unique")
        require(set(order) == selected_experiences, "Experience order must exactly cover selected experiences")
        require(data.get("full_recalculation") is True, "Recompute the full match table after evidence updates")
        require(
            data.get("backfill_state") in ("未触发", "待回补", "已重新计算", "用户确认无更多经历后已重新计算"),
            "Invalid backfill state",
        )

    if data.get("kind") == "match":
        require(not data.get("header") and not data.get("sections"), "Match record must not contain resume content")
        require(nonempty(data.get("review_notes")), "Review notes required")
        return data

    header = data.get("header")
    require(isinstance(header, list), "Header must be list")
    lines = [(line, False) for line in header]
    require(isinstance(data.get("sections"), list), "Sections required")
    for section in data["sections"]:
        require(nonempty(section.get("title")) and isinstance(section.get("lines"), list), "Invalid section")
        if data.get("kind") == "tailored":
            require(section.get("section_type") in ("education", "experience", "skills", "other"), "Tailored section type required")
        lines += [(line, True, section.get("section_type")) for line in section["lines"]]
    lines = [(line, False, None) for line in header] + [item for item in lines if item[1]]
    require(bool(lines), "Resume content required")
    experience_positions = []
    for line, is_body, section_type in lines:
        require(nonempty(line.get("text")), "Empty resume line")
        refs = line.get("evidence_ids")
        require(
            isinstance(refs, list) and refs and all(ref in evidence for ref in refs),
            "Every line needs existing evidence references",
        )
        if data.get("kind") == "tailored":
            require(
                not re.search(r"弱匹配|匹配状态|直接证据|代理证据|证据等级|核心缺口|当前材料未提供|不能\s*final|投递前.*谨慎", line["text"], re.I),
                "Backend diagnosis cannot enter resume body",
            )
        if data["stage"] == "final":
            require(
                all(evidence[ref]["status"] in ("材料已支持", "用户已确认") for ref in refs),
                "Unresolved evidence cannot enter final resume",
            )
            require(
                not re.search(r"【[^】]*】|\[(?:TODO|TBD|metric|占位符)\]|待填写|待补充", line["text"], re.I),
                "Placeholder in final resume",
            )
        if data.get("kind") == "tailored" and is_body:
            require(line.get("line_type") in ("identity", "evidence"), "Tailored line type required")
            if line["line_type"] == "identity":
                require(not line.get("task_ids"), "Identity line must not claim a JD task")
                continue
            task_refs = line.get("task_ids")
            require(
                isinstance(task_refs, list)
                and task_refs
                and all(ref in task_ids for ref in task_refs),
                "Every evidence line must support a current JD task cluster",
            )
            task_selected_evidence = {
                ref
                for item in data["match_matrix"]
                if item["selected_for_resume"] and item["task_id"] in task_refs
                for ref in item["evidence_ids"]
            }
            require(all(ref in task_selected_evidence for ref in refs), "Resume line uses evidence not selected for its task")
            require(all(ref in data["allowed_evidence_ids"] for ref in refs), "Resume line uses evidence outside the allowed evidence set")
            require(line.get("experience_id") in data["experience_order"], "Resume evidence line needs an allowed experience id")
            experience_positions.append(data["experience_order"].index(line["experience_id"]))
            if section_type == "skills":
                require(line.get("skill_type") in ("软件", "工具", "平台", "方法", "证书", "语言", "可操作技能"), "Skills must be verifiable")
                require(not re.search(r"能力|意识|思维|素养$", line["text"]), "Generic self-evaluation cannot be a skill")

    if data.get("kind") == "tailored":
        require(experience_positions == sorted(experience_positions), "Resume experience order must follow the saved JD proof order")

    if data.get("kind") == "tailored":
        require(nonempty(data.get("source_match_version")), "Tailored version must cite a saved match record")
        require(nonempty(data.get("source_master_version")), "Tailored version must cite the latest asset master")
        require(data.get("facts_reloaded") is True, "Reload current facts before generating a resume")
        require(data.get("allowed_set_reloaded") is True, "Reload the current allowed evidence set before generating a resume")
        require(isinstance(data.get("dates_required"), bool), "Resume date requirement decision required")
        require(isinstance(data.get("dates_batch_checked"), bool), "Batch date check state required")
        if data["dates_required"]:
            require(data["dates_batch_checked"] is True, "Batch-check all selected experience dates before generation")
            require(
                all(experiences[experience_id]["start_date"] != "未知" and experiences[experience_id]["end_date"] != "未知" for experience_id in data["experience_order"]),
                "Required resume dates cannot remain unknown",
            )
        require("version_mode" not in data, "Match strength belongs in the JD match record, not the resume")
        unresolved_core = [
            item for item in data["match_matrix"]
            if item["core"] and item["status"] != "已有直接证据"
        ]
        require(
            all(item["gap_resolution"] != "待回补" for item in unresolved_core),
            "Core evidence gap still needs targeted backfill; do not create a resume",
        )
        for field in ("new_evidence", "removed_or_compressed", "added_or_expanded"):
            require(isinstance(data.get(field), list), "Tailored changelog fields required")
        require(isinstance(data.get("hard_gate_blockers"), list), "Hard-gate blocker list required")

    if data["stage"] == "final":
        require(data["kind"] == "tailored", "Final application resume requires a specific JD")
        require(data.get("contact_checked") is True, "Name/contact need review")
        checks = data.get("checks")
        require(isinstance(checks, dict), "Review checks required")
        require(checks.get("真实性") == "通过当前检查", "Authenticity check required")
        require(checks.get("岗位相关性") in ("已检查", "通过当前检查"), "Relevance must be checked")
        require(checks.get("可读性") == "通过当前检查", "Readability check required")
        require(data.get("blockers") == [], "Resolve blockers before final")
        require(data.get("hard_gate_blockers") == [], "Explicit hard-gate blockers prevent final")
    require(nonempty(data.get("review_notes")), "Review notes required")
    return data


def plain(data):
    parts = [item["text"] for item in data["header"]]
    for section in data["sections"]:
        parts += ["", section["title"]] + [item["text"] for item in section["lines"]]
    return "\n".join(parts) + "\n"


def render_profile(data):
    labels = {
        "alias": "称呼",
        "education": "学历",
        "school": "学校",
        "major": "专业",
        "graduation": "毕业时间",
        "search_stage": "求职阶段",
        "target_directions": "目标方向",
    }
    rows = ["# 个人资料", ""]
    for key, label in labels.items():
        value = data.get("profile", {}).get(key)
        if isinstance(value, list):
            value = "、".join(str(x) for x in value)
        if value not in (None, "", []):
            rows.append(f"- {label}：{value}")
    if data.get("workflow_state"):
        rows.append("- 工作流状态：" + data["workflow_state"])
    return "\n".join(rows) + "\n"


def render_experience(data):
    parts = ["# 真实经历母版", ""]
    parts += ["## 结构化经历索引", ""]
    for item in data["experiences"]:
        parts += [
            "### " + item["experience_id"] + "｜" + item["name"],
            "",
            "- 来源类型：" + item["source_type"],
            "- 完成状态：" + item["execution_status"],
            "- 开始时间：" + item["start_date"],
            "- 结束时间：" + item["end_date"],
            "- 持续时间：" + item["duration"],
            "- 证据：" + "、".join(item["evidence_ids"]),
            "",
        ]
    for section in data["sections"]:
        parts += ["## " + section["title"], ""]
        for item in section["lines"]:
            parts.append("- " + item["text"] + "〔证据：" + "、".join(item["evidence_ids"]) + "〕")
        parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def render_evidence(data):
    rows = ["# 证据链", ""]
    for item in data["evidence"]:
        rows += [
            "## " + item["id"],
            "",
            "- 状态：" + item["status"],
            "- 事实：" + item["fact"],
            "- 来源：" + item["source"],
            "- 归属：" + item["ownership"],
            "- 位置：" + (item.get("locator") or "未单独标注"),
            "",
        ]
    return "\n".join(rows).rstrip() + "\n"


def render_signals(data):
    rows = ["# 岗位信号", ""]
    if not data["job_signals"]:
        rows.append("尚未从现有证据中提取出岗位信号。")
    for item in data["job_signals"]:
        rows += [
            "## " + item["name"],
            "",
            "- 支持范围：" + item["scope"],
            "- 证据：" + "、".join(item["evidence_ids"]),
            "",
        ]
    return "\n".join(rows).rstrip() + "\n"


def render_jd(data):
    return (
        "# " + data["company"] + "｜" + data["target"] + "\n\n"
        + "- JD编号：" + data["jd_id"] + "\n"
        + "- 保存时间：" + datetime.now(timezone.utc).isoformat() + "\n\n"
        + "## JD原文\n\n" + data["jd_text"].strip() + "\n"
    )


def render_match(data):
    rows = ["# JD 匹配记录", "", "- JD编号：" + data["jd_id"], "", "## 核心任务簇", ""]
    for cluster in data["task_clusters"]:
        rows += [
            "### " + cluster["task_id"] + "｜" + cluster["name"],
            "",
            "- 任务对象：" + cluster["object"],
            "- 反复动作：" + cluster["action"],
            "- 交付物或结果：" + cluster["deliverable"],
            "- 期望证据：" + cluster["evidence_expected"],
            "",
        ]
    rows += [
        "## 当前简历允许使用的证据集合",
        "",
        "- 证据白名单：" + ("、".join(data["allowed_evidence_ids"]) or "无"),
        "- 经历顺序：" + ("、".join(data["experience_order"]) or "无"),
        "- 回补状态：" + data["backfill_state"],
        "- 整表重算：已完成",
        "",
        "## 证据映射",
        "",
    ]
    for item in data["match_matrix"]:
        rows += [
            "## " + item["requirement_id"] + "｜" + item["requirement"],
            "",
            "- 核心要求：" + ("是" if item["core"] else "否"),
            "- 所属任务：" + item["task_id"],
            "- 经历编号：" + (item.get("experience_id") or "无"),
            "- 岗位信号：" + (item.get("signal") or "无"),
            "- 证据：" + ("、".join(item["evidence_ids"]) or "无"),
            "- 匹配状态：" + item["status"],
            "- 证据等级：" + item["evidence_grade"],
            "- 岗位相关性：" + item["relevance"],
            "- 个人贡献：" + item["contribution_clarity"],
            "- 结果或交付物：" + item["result_strength"],
            "- 入选简历：" + ("是" if item["selected_for_resume"] else "否"),
            "- 回补状态：" + item["gap_resolution"],
            "- 处理：" + item["action"],
            "",
        ]
    return "\n".join(rows).rstrip() + "\n"


def document_html(data):
    esc = html.escape
    body = "".join("<p>" + esc(item["text"]) + "</p>" for item in data["header"])
    for section in data["sections"]:
        body += "<h2>" + esc(section["title"]) + "</h2>"
        body += "".join("<p>" + esc(item["text"]) + "</p>" for item in section["lines"])
    return '<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>简历</title><style>@page{size:A4;margin:18mm}body{font:11pt/1.5 Arial,Calibri,"Noto Sans CJK SC","Microsoft YaHei",sans-serif;max-width:180mm;margin:auto;color:#111}h2{font-size:13pt;border-bottom:1px solid #aaa;break-after:avoid}p{white-space:pre-wrap;overflow-wrap:anywhere;margin:5pt 0}</style><body>' + body + "</body></html>"


def load(case, version_id):
    case = require_case(case)
    require(bool(re.fullmatch(r"[A-Za-z0-9_-]+", version_id)), "Invalid version ID")
    path = case / ".records" / (version_id + ".json")
    require(path.is_file(), "Version not found")
    return json.loads(path.read_text())


def records(case):
    case = require_case(case)
    result = []
    for path in sorted((case / ".records").glob("*.json")):
        result.append(json.loads(path.read_text()))
    return result


def version_list(case):
    case = require_case(case)
    labels = {}
    for path in sorted((case / ".labels").glob("*.json")):
        event = json.loads(path.read_text())
        labels[event["version_id"]] = event["name"]
    rows = []
    for record in records(case):
        data = record["data"]
        rows.append(
            {
                "version_id": record["version_id"],
                "name": labels.get(record["version_id"], record["version_name"]),
                "company": data.get("company"),
                "target": data.get("target"),
                "jd_id": data.get("jd_id"),
                "branch_version": record["branch_version"],
                "kind": data["kind"],
                "stage": data["stage"],
                "parent_version": data.get("parent_version"),
                "created_at": record["created_at"],
                "change_reason": data["change_reason"],
            }
        )
    return rows


def append_changelog(case, record):
    data = record["data"]
    gaps = data.get("remaining_gaps") or []
    gap_text = "；".join(gaps) if gaps else "无"
    parent = data.get("parent_version") or record.get("master_version") or "无，首次记录"
    if data["kind"] == "tailored":
        new_evidence = "；".join(data["new_evidence"]) if data["new_evidence"] else "无"
        removed = "；".join(data["removed_or_compressed"]) if data["removed_or_compressed"] else "无"
        added = "；".join(data["added_or_expanded"]) if data["added_or_expanded"] else "无"
        entry = (
            "\n## " + record["version_name"] + "\n\n"
            + "- 版本：" + record["version_name"] + "\n"
            + "- 对应JD：" + data["jd_id"] + "｜" + data["company"] + "｜" + data["target"] + "\n"
            + "- 基于哪个版本：" + parent + "\n"
            + "- 本轮新增证据：" + new_evidence + "\n"
            + "- 删除／压缩了什么：" + removed + "\n"
            + "- 新增／展开了什么：" + added + "\n"
            + "- 修改原因：" + data["change_reason"] + "\n"
            + "- 仍未解决的核心缺口：" + gap_text + "\n"
            + "- 版本ID：" + record["version_id"] + "\n"
        )
    else:
        changes = "；".join(data["changes"])
        corresponding = data.get("jd_id") or "素材母版"
        entry = (
            "\n## " + record["version_name"] + "\n\n"
            + "- 版本：" + record["version_name"] + "\n"
            + "- 对应JD：" + corresponding + "\n"
            + "- 基于哪个版本：" + parent + "\n"
            + "- 修改原因：" + data["change_reason"] + "\n"
            + "- 本次变化：" + changes + "\n"
            + "- 仍未解决的核心缺口：" + gap_text + "\n"
            + "- 版本ID：" + record["version_id"] + "\n"
        )
    with (case / "changelog.md").open("a") as handle:
        handle.write(entry)


def update_master_assets(case, data):
    atomic_write(case / "profile.md", render_profile(data))
    atomic_write(case / "experience-master.md", render_experience(data))
    atomic_write(case / "evidence.md", render_evidence(data))
    atomic_write(case / "job-signals.md", render_signals(data))


def ensure_jd(case, data):
    jd_id = safe_name(data["jd_id"])
    company = safe_name(data["company"])
    target = safe_name(data["target"])
    jd_path = case / "jds" / (jd_id + "-" + company + "-" + target + ".md")
    if jd_path.exists():
        old_body = jd_path.read_text().split("## JD原文\n\n", 1)[-1].strip()
        require(old_body == data["jd_text"].strip(), "Existing JD id has different text; create a new JD id")
    else:
        jd_path.write_text(render_jd(data))
    return jd_path


def save(case, data):
    case = require_case(case)
    validate(data)
    existing = records(case)
    parent = data.get("parent_version")
    parent_record = load(case, parent) if parent else None
    master_records = [item for item in existing if item["data"]["kind"] == "master"]
    master_version = None

    if data["kind"] == "master":
        if master_records:
            require(parent_record is not None, "Existing master requires parent_version")
        require(not parent_record or parent_record["data"]["kind"] == "master", "Master parent must be a master version")
        branch_version = 1 + max(
            [item["branch_version"] for item in master_records] or [0]
        )
        version_name = "素材母版-v" + str(branch_version)
        artifact_path = None
        update_master_assets(case, data)
    elif data["kind"] == "match":
        require(master_records, "Create and save the evidence master before a tailored branch")
        latest_master = max(master_records, key=lambda item: item["branch_version"])
        master_version = latest_master["version_id"]
        require(data.get("source_master_version") == master_version, "Match must cite the latest asset master")
        master_evidence = {item["id"] for item in latest_master["data"]["evidence"]}
        match_evidence = {item["id"] for item in data["evidence"]}
        require(
            match_evidence.issubset(master_evidence),
            "Match evidence must already exist in the latest master",
        )
        branch = [item for item in existing if item["data"].get("kind") == "match" and item["data"].get("jd_id") == data["jd_id"]]
        if branch:
            require(parent_record is not None, "Existing JD branch requires parent_version")
        if parent_record is not None:
            require(
                parent_record["data"].get("kind") == "match" and parent_record["data"].get("jd_id") == data["jd_id"],
                "Match parent must belong to the same JD",
            )
        branch_version = 1 + max([item["branch_version"] for item in branch] or [0])
        company = safe_name(data["company"])
        target = safe_name(data["target"])
        jd_id = safe_name(data["jd_id"])
        version_name = company + "-" + target + "-" + jd_id + "-匹配-v" + str(branch_version)
        ensure_jd(case, data)
        artifact_path = case / "jds" / (version_name + ".md")
        require(not artifact_path.exists(), "Match record path already exists")
        artifact_path.write_text(render_match(data))
    else:
        require(master_records, "Create and save the evidence master before a tailored branch")
        latest_master = max(master_records, key=lambda item: item["branch_version"])
        master_version = latest_master["version_id"]
        master_evidence = {item["id"] for item in latest_master["data"]["evidence"]}
        tailored_evidence = {item["id"] for item in data["evidence"]}
        require(tailored_evidence.issubset(master_evidence), "Tailored evidence must already exist in the latest master")
        match_record = load(case, data["source_match_version"])
        require(match_record["data"]["kind"] == "match", "source_match_version must be a match record")
        require(match_record["data"]["jd_id"] == data["jd_id"], "Match record must belong to the same JD")
        require(data["source_master_version"] == latest_master["version_id"], "Reload and cite the latest asset master")
        require(match_record["data"]["source_master_version"] == latest_master["version_id"], "Saved match is stale; rebuild it from the latest master")
        require(match_record["data"]["task_clusters"] == data["task_clusters"], "Resume must use the saved JD task model")
        require(match_record["data"]["match_matrix"] == data["match_matrix"], "Resume must use the saved match matrix")
        require(match_record["data"]["allowed_evidence_ids"] == data["allowed_evidence_ids"], "Resume must use the saved allowed evidence set")
        require(match_record["data"]["experience_order"] == data["experience_order"], "Resume must use the saved JD proof order")
        require(match_record["data"]["full_recalculation"] is True, "Saved match must be fully recalculated")
        ensure_jd(case, data)
        branch = [item for item in existing if item["data"].get("kind") == "tailored" and item["data"].get("jd_id") == data["jd_id"]]
        if branch:
            require(parent_record is not None, "Existing resume branch requires parent_version")
        if parent_record is not None:
            require(parent_record["data"].get("kind") == "tailored" and parent_record["data"].get("jd_id") == data["jd_id"], "Parent version must belong to the same resume branch")
        branch_version = 1 + max([item["branch_version"] for item in branch] or [0])
        company = safe_name(data["company"])
        target = safe_name(data["target"])
        jd_id = safe_name(data["jd_id"])
        version_name = company + "-" + target + "-" + jd_id + "-v" + str(branch_version)
        artifact_path = case / "resumes" / (version_name + ".md")
        require(not artifact_path.exists(), "Resume version path already exists")
        artifact_path.write_text(plain(data))

    version_id = timestamp_id()
    record = {
        "version_id": version_id,
        "version_name": version_name,
        "branch_version": branch_version,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "artifact_path": str(artifact_path) if artifact_path else None,
        "master_version": master_version,
        "data": data,
    }
    record_path = case / ".records" / (version_id + ".json")
    with record_path.open("x") as handle:
        json.dump(record, handle, ensure_ascii=False, indent=2)
    require(load(case, version_id) == record, "Readback failed")
    append_changelog(case, record)
    return {
        "case_dir": str(case),
        "version_id": version_id,
        "name": version_name,
        "parent_version": parent,
        "artifact": str(artifact_path) if artifact_path else None,
        "resume": str(artifact_path) if data["kind"] == "tailored" else None,
        "changelog": str(case / "changelog.md"),
    }


def rename(case, version_id, name):
    case = require_case(case)
    load(case, version_id)
    require(nonempty(name), "Version name required")
    event_id = timestamp_id()
    with (case / ".labels" / (event_id + ".json")).open("x") as handle:
        json.dump({"version_id": version_id, "name": name}, handle, ensure_ascii=False)
    return {"version_id": version_id, "name": name}


def rollback(case, version_id):
    old = load(case, version_id)
    data = json.loads(json.dumps(old["data"], ensure_ascii=False))
    data["parent_version"] = version_id
    data["restored_from"] = version_id
    data["stage"] = "draft"
    data["checks"] = {key: "待补" for key in ("真实性", "岗位相关性", "可读性")}
    data["blockers"] = ["恢复后需核对最新事实、证据与目标JD"]
    data["change_reason"] = "从历史版本恢复为新草稿，保留原版本和后续版本。"
    data["changes"] = ["恢复历史正文，等待按当前事实和JD复核"]
    return save(case, data)


def export(case, version_id, fmt):
    case = require_case(case)
    record = load(case, version_id)
    data = record["data"]
    validate(data)
    require(data["kind"] == "tailored", "Export a tailored resume version, not the master")
    require(fmt in ("html", "docx", "pdf"), "Invalid format")
    destination = case / "exports" / (record["version_name"] + "-" + uuid.uuid4().hex[:8])
    destination.mkdir(parents=True)
    base = record["version_name"] if data["stage"] == "final" else record["version_name"] + "-草稿"
    output = destination / (base + "." + fmt)
    warning = None
    if fmt == "html":
        output.write_text(document_html(data))
    elif fmt == "docx":
        try:
            from docx import Document
        except ImportError as error:
            raise RuntimeError("DOCX需要python-docx；可以先导出HTML。") from error
        from docx.shared import Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        for style_name in ("Normal", "Heading 1"):
            style = document.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(11 if style_name == "Normal" else 13)
            props = style.element.get_or_add_rPr()
            fonts = props.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                props.append(fonts)
            fonts.set(qn("w:eastAsia"), "宋体")
        for item in data["header"]:
            document.add_paragraph(item["text"])
        for section in data["sections"]:
            document.add_heading(section["title"], level=1)
            for item in section["lines"]:
                document.add_paragraph(item["text"])
        document.save(output)
        readback = "\n".join(item.text for item in Document(output).paragraphs).strip()
        require(readback == plain(data).replace("\n\n", "\n").strip(), "DOCX readback mismatch")
    else:
        try:
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.lib.pagesizes import A4
        except ImportError:
            output = destination / (base + ".html")
            output.write_text(document_html(data))
            warning = "PDF依赖缺失，实际生成HTML，可打印为PDF；没有生成PDF。"
        else:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=11, leading=16, wordWrap="CJK")
            story = []
            for text in plain(data).splitlines():
                story.extend([Paragraph(html.escape(text) or " ", style), Spacer(1, 4)])
            SimpleDocTemplate(
                str(output), pagesize=A4, rightMargin=45, leftMargin=45, topMargin=45, bottomMargin=45
            ).build(story)
    require(output.exists() and output.stat().st_size > 0, "Empty export")
    return {
        "file": str(output),
        "actual_format": output.suffix[1:],
        "warning": warning,
        "visual_review": "尚需检查版面",
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", type=Path, default=Path.home() / "Documents" / "resume-cases")
    parser.add_argument("--case", type=Path)
    commands = parser.add_subparsers(dest="cmd", required=True)

    command = commands.add_parser("init-case")
    command.add_argument("--alias", required=True)
    command.add_argument("--date")

    command = commands.add_parser("save")
    command.add_argument("--payload", type=Path, required=True)

    command = commands.add_parser("export")
    command.add_argument("--id", required=True)
    command.add_argument("--format", choices=["html", "docx", "pdf"], required=True)

    command = commands.add_parser("compare")
    command.add_argument("--old", required=True)
    command.add_argument("--new", required=True)

    command = commands.add_parser("rename")
    command.add_argument("--id", required=True)
    command.add_argument("--name", required=True)

    command = commands.add_parser("rollback")
    command.add_argument("--id", required=True)
    commands.add_parser("list")

    args = parser.parse_args()
    if args.cmd == "init-case":
        result = init_case(args.base, args.alias, args.date)
    else:
        require(args.case is not None, "--case is required")
        case = args.case
        if args.cmd == "save":
            result = save(case, json.loads(args.payload.read_text()))
        elif args.cmd == "rename":
            result = rename(case, args.id, args.name)
        elif args.cmd == "rollback":
            result = rollback(case, args.id)
        elif args.cmd == "export":
            result = export(case, args.id, args.format)
        elif args.cmd == "compare":
            result = "".join(
                difflib.unified_diff(
                    plain(load(case, args.old)["data"]).splitlines(True),
                    plain(load(case, args.new)["data"]).splitlines(True),
                    fromfile=args.old,
                    tofile=args.new,
                )
            )
        else:
            result = version_list(case)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
